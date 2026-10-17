# -*- coding: utf-8 -*-
"""Acrescenta ao livro de gráficos a folha de testes emparelhados: Friedman,
Q de Cochran e McNemar com fórmulas vivas, e o quadro das comparações par a par
de Wilcoxon com correção de Holm."""
import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from graficos_nativos import HDR2, BR, titulo

PICK='Teste_pick_TM12_final_CORRIGIDO.xlsx'
OBST='Testes_Obstaculos_TM12_final_CORRIGIDO.xlsx'
TEMPO='Testes_de_tempo_e_estados_TM12_AUDITADO.xlsx'
PL=['RRTConnect','RRTstar','PRMstar']
MET=[('Tempo total (s)',5),('Comprimento (m)',6),('Tempo de planeamento (s)',3)]

def ler(f,nc=11,folha='Dados Análise'):
    ws=openpyxl.load_workbook(f,data_only=True)[folha]
    return [[ws.cell(r,c).value for c in range(1,nc+1)] for r in range(2,ws.max_row+1)]

def num(v): return isinstance(v,(int,float))

wb=openpyxl.load_workbook('Graficos_Capitulo5_EDITAVEIS.xlsx')
ws=wb.create_sheet('Testes emparelhados',2)
ws.sheet_view.showGridLines=False
titulo(ws,'A1','TESTES EMPARELHADOS — FRIEDMAN, WILCOXON COM HOLM, Q DE COCHRAN E McNEMAR')
def nota(cel,t): ws[cel]=t; ws[cel].font=Font(italic=True,size=9,color='595959')
def sub(cel,t): ws[cel]=t; ws[cel].font=Font(bold=True,size=10.5,color='1F3864')
def hdr(r,c,txts,larg=None):
    for j,t in enumerate(txts):
        cl=ws.cell(r,c+j,t); cl.font=Font(bold=True,size=9.5); cl.fill=HDR2
        cl.alignment=Alignment(horizontal='center',wrap_text=True); cl.border=BR
    if larg:
        for j,w in enumerate(larg): ws.column_dimensions[get_column_letter(c+j)].width=w

nota('A2','O desenho é emparelhado: cada movimento foi executado pelos três planeadores e cada ciclo pelas '
          'duas arquiteturas. Os testes para grupos independentes não são aplicáveis. Tudo o que está abaixo é fórmula.')
ws.column_dimensions['A'].width=16; ws.column_dimensions['B'].width=10

def blocos(fich, chave):
    L=[r for r in ler(fich) if r[2] in PL]
    d={p:{} for p in PL}
    for r in L:
        k=chave(r)
        d[r[2]][k]=r if all(num(r[i]) for i in (3,5,6)) else None
    todas=sorted(set().union(*[set(x) for x in d.values()]), key=lambda z: str(z))
    completos=[k for k in todas if all(d[p].get(k) for p in PL)]
    return d,todas,completos

linha=4
for fich,nome,chave in [(PICK,'CAMPANHA 1 — ciclo completo, sem obstáculos',lambda r:(r[0],r[1])),
                        (OBST,'CAMPANHA 2 — movimento isolado, com obstáculos',lambda r:(r[1],r[0]))]:
    d,todas,completos=blocos(fich,chave)
    sub(f'A{linha}',nome); linha+=1
    nota(f'A{linha}',f'Blocos totais: {len(todas)}. Blocos completos (concluídos pelos três planeadores): {len(completos)}. '
                     f'As métricas contínuas usam os blocos completos; a taxa de sucesso usa todos.')
    linha+=2
    r0=linha
    cabec=['Bloco']
    for mnome,_ in MET: cabec+= [f'{p} — {mnome}' for p in PL]+[f'posto {p}' for p in PL]
    hdr(r0,1,cabec,[16]+[13]*18)
    for i,k in enumerate(completos):
        r=r0+1+i
        ws.cell(r,1,str(k)).font=Font(size=9)
        for m,(mnome,idx) in enumerate(MET):
            base=2+m*6
            for j,p in enumerate(PL):
                cl=ws.cell(r,base+j,float(d[p][k][idx])); cl.number_format='0.0000'; cl.font=Font(size=9)
            a=get_column_letter(base); z=get_column_letter(base+2)
            for j in range(3):
                col=get_column_letter(base+j)
                cl=ws.cell(r,base+3+j,f'=RANK({col}{r},${a}{r}:${z}{r},1)')
                cl.number_format='0'; cl.font=Font(size=9)
    fim=r0+len(completos)
    linha=fim+2
    hdr(linha,1,['Métrica','n blocos','ΣR RRTConnect','ΣR RRTstar','ΣR PRMstar','χ² de Friedman','gl','p'],
        [26,10,15,15,15,16,6,12])
    for m,(mnome,_) in enumerate(MET):
        r=linha+1+m; base=2+m*6+3
        ws.cell(r,1,mnome).font=Font(size=10)
        ws.cell(r,2,len(completos)).number_format='0'
        for j in range(3):
            col=get_column_letter(base+j)
            ws.cell(r,3+j,f'=SUM({col}{r0+1}:{col}{fim})').number_format='0'
        n=len(completos)
        ws.cell(r,6,f'=12/({n}*3*4)*(C{r}^2+D{r}^2+E{r}^2)-3*{n}*4').number_format='0.000'
        ws.cell(r,7,2).number_format='0'
        ws.cell(r,8,f'=CHIDIST(F{r},G{r})').number_format='0.0000'
        for c in range(1,9): ws.cell(r,c).border=BR
    # Q de Cochran sobre todos os blocos
    linha=linha+5
    suc={p:sum(1 for k in todas if d[p].get(k)) for p in PL}
    N=len(todas)
    Li={}
    for k in todas: Li[k]=sum(1 for p in PL if d[p].get(k))
    somaLi2=sum(v*v for v in Li.values()); total=sum(suc.values())
    hdr(linha,1,['Taxa de sucesso','Sucessos RRTConnect','Sucessos RRTstar','Sucessos PRMstar','Blocos','Q de Cochran','gl','p'],
        [26,17,15,15,10,14,6,12])
    r=linha+1
    ws.cell(r,1,'Blocos concluídos').font=Font(size=10)
    for j,p in enumerate(PL): ws.cell(r,2+j,suc[p]).number_format='0'
    ws.cell(r,5,N).number_format='0'
    ws.cell(r,6,f'=2*(3*(B{r}^2+C{r}^2+D{r}^2)-{total}^2)/(3*{total}-{somaLi2})').number_format='0.000'
    ws.cell(r,7,2).number_format='0'
    ws.cell(r,8,f'=CHIDIST(F{r},G{r})').number_format='0.0000'
    for c in range(1,9): ws.cell(r,c).border=BR
    linha=r+3

# ---- Wilcoxon par a par com Holm (valores do script) ----
sub(f'A{linha}','COMPARAÇÕES PAR A PAR — Wilcoxon emparelhado com correção de Holm-Bonferroni')
linha+=1
nota(f'A{linha}','Calculado por testes_emparelhados.py (scipy). O Excel não dispõe de função para o teste de '
                 'Wilcoxon emparelhado; os valores são reproduzíveis correndo o script sobre os mesmos ficheiros.')
linha+=2
hdr(linha,1,['Campanha','Métrica','Par comparado','p de Wilcoxon','p corrigido por Holm','Significativo a 5 %'],
    [22,24,30,16,20,18])
DADOS=[('Sem obstáculos','Tempo total','RRTConnect vs. RRTstar',0.001197,0.003592,'Sim'),
       ('','','RRTConnect vs. PRMstar',0.010630,0.021260,'Sim'),
       ('','','RRTstar vs. PRMstar',0.441200,0.441200,'Não'),
       ('','Comprimento','RRTConnect vs. RRTstar',0.0003254,0.0006508,'Sim'),
       ('','','RRTConnect vs. PRMstar',0.00002734,0.00008202,'Sim'),
       ('','','RRTstar vs. PRMstar',0.036820,0.036820,'Sim'),
       ('','Tempo de planeamento','RRTConnect vs. RRTstar',0.002889,0.005779,'Sim'),
       ('','','RRTConnect vs. PRMstar',0.379200,0.379200,'Não'),
       ('','','RRTstar vs. PRMstar',0.0002504,0.0007513,'Sim'),
       ('Com obstáculos','—','Friedman não significativo; não se fazem comparações par a par','','','')]
for i,d_ in enumerate(DADOS):
    r=linha+1+i
    for j,v in enumerate(d_):
        cl=ws.cell(r,1+j,v); cl.font=Font(size=9.5); cl.border=BR
        if isinstance(v,float): cl.number_format='0.000000'
linha=linha+len(DADOS)+3

# ---- McNemar: temporal vs por respostas ----
sub(f'A{linha}','TAXA DE SUCESSO DO CICLO — McNemar exato (substitui o teste de Fisher)')
linha+=1
# contagens lidas da base auditada (coluna estado_inclusao = «incluído»)
_wa=openpyxl.load_workbook(TEMPO,data_only=True)['Dados Análise']
_aud=[[_wa.cell(r,c).value for c in (8,9,58)] for r in range(2,_wa.max_row+1)]
_aud=[x for x in _aud if str(x[2]).strip().lower()=='incluído']
nota(f'A{linha}','Cada repetição foi executada pelas duas estratégias nas mesmas condições, pelo que os '
                 'resultados são emparelhados e o teste de Fisher, que pressupõe amostras independentes, '
                 f'não é aplicável. A tabela assenta nas {len(_aud)} condições com registo primário das duas '
                 'estratégias, apuradas na auditoria de proveniência.')
linha+=2
hdr(linha,1,['','Sucesso por respostas','Falha por respostas','Total'],[26,20,20,12])
_n11=sum(1 for a,b,_ in _aud if a==1 and b==1); _n10=sum(1 for a,b,_ in _aud if a==1 and b!=1)
_n01=sum(1 for a,b,_ in _aud if a!=1 and b==1); _n00=sum(1 for a,b,_ in _aud if a!=1 and b!=1)
vals=[['Sucesso temporal',_n11,_n10,None],['Falha temporal',_n01,_n00,None],['Total',None,None,None]]
for i,v in enumerate(vals):
    r=linha+1+i
    ws.cell(r,1,v[0]).font=Font(size=10)
    if v[1] is not None: ws.cell(r,2,v[1]).number_format='0'
    if v[2] is not None: ws.cell(r,3,v[2]).number_format='0'
    for c in range(1,5): ws.cell(r,c).border=BR
b=linha+1
ws.cell(b,4,f'=B{b}+C{b}').number_format='0'
ws.cell(b+1,4,f'=B{b+1}+C{b+1}').number_format='0'
ws.cell(b+2,2,f'=B{b}+B{b+1}').number_format='0'
ws.cell(b+2,3,f'=C{b}+C{b+1}').number_format='0'
ws.cell(b+2,4,f'=D{b}+D{b+1}').number_format='0'
r=b+4
r_disc, r_menor = r, r + 1
for rot,f in [('Pares discordantes',f'=C{b}+B{b+1}'),
              ('Menor das duas contagens',f'=MIN(C{b},B{b+1})'),
              ('p exato bilateral (binomial)',f'=MIN(1,2*BINOMDIST(B{r_menor},B{r_disc},0.5,TRUE))')]:
    ws.cell(r,1,rot).font=Font(size=10); cl=ws.cell(r,2,f)
    cl.number_format='0' if 'p ' not in rot else '0.0000'
    for c in (1,2): ws.cell(r,c).border=BR
    r+=1

wb.calculation.fullCalcOnLoad=True
wb.save('Graficos_Capitulo5_EDITAVEIS.xlsx')
print('folha «Testes emparelhados» criada. folhas:',len(wb.sheetnames))
