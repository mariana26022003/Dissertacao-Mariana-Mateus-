# -*- coding: utf-8 -*-
"""Completa o livro de gráficos: barras de erro nos gráficos que as devem ter e
os três gráficos de planeamento que faltavam para o Capítulo 5."""
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.chart.data_source import NumDataSource, NumRef
from graficos_nativos import C1, C2, C3, HDR2, BR, titulo
from openpyxl.styles import Font, Alignment

wb = openpyxl.load_workbook('Graficos_Capitulo5_EDITAVEIS.xlsx')

def erro(ch, refs):
    for k, s in enumerate(ch.series):
        if k < len(refs) and refs[k]:
            s.errBars = ErrorBars(errDir='y', errValType='cust', errBarType='both',
                                  plus=NumDataSource(NumRef(f=refs[k])),
                                  minus=NumDataSource(NumRef(f=refs[k])))

def acha(ws, frag):
    for ch in ws._charts:
        t = ch.title.tx.rich.p[0].r[0].t if ch.title else ''
        if frag in t: return ch
    raise KeyError(frag)

# --- IC 95 % no gráfico por cenário -----------------------------------
ws = wb['Sincronização']
erro(acha(ws, 'Tempo médio de ciclo por cenário'),
     ['Sincronização!$E$5:$E$12', 'Sincronização!$F$5:$F$12'])

# --- desvio-padrão nos gráficos dos planeadores -----------------------
pl = wb['Planeadores']
erro(acha(pl, 'Tempo total médio por planeador — ciclo completo'), ['Planeadores!$G$5:$G$7'])
erro(acha(pl, 'Tempo total médio por planeador — movimento'), ['Planeadores!$G$12:$G$14'])
erro(acha(pl, 'Comprimento cartesiano médio da trajetória — sem'), ['Planeadores!$I$5:$I$7'])
erro(acha(pl, 'Comprimento cartesiano médio da trajetória — com'), ['Planeadores!$I$12:$I$14'])

# --- tabela e gráfico do número de pontos ------------------------------
pl['A66'] = 'Discretização da trajetória (número médio de pontos)'
pl['A66'].font = Font(bold=True, size=11, color='1F3864')
for j, h in enumerate(['Planeador', 'Sem obstáculos', 'Com obstáculos']):
    c = pl.cell(67, 1 + j, h); c.font = Font(bold=True, size=10); c.fill = HDR2
    c.alignment = Alignment(horizontal='center', wrap_text=True); c.border = BR
for i in range(3):
    pl.cell(68 + i, 1, f'=A{5+i}').font = Font(size=10)
    pl.cell(68 + i, 2, f'=J{5+i}').number_format = '0.0'
    pl.cell(68 + i, 3, f'=J{12+i}').number_format = '0.0'
    for c in range(1, 4): pl.cell(68 + i, c).border = BR

def barras(ws, anc, tit, ex, ey, dados, cats, cores, err=None, w=15, h=9, leg=True):
    ch = BarChart(); ch.type = 'col'
    ch.title = tit; ch.x_axis.title = ex; ch.y_axis.title = ey
    ch.add_data(dados, titles_from_data=leg)
    ch.set_categories(cats)
    ch.width, ch.height = w, h
    ch.gapWidth = 70
    for k, s in enumerate(ch.series):
        s.graphicalProperties.solidFill = cores[k % len(cores)]
        s.graphicalProperties.line.solidFill = cores[k % len(cores)]
    if err: erro(ch, err)
    if not leg: ch.legend = None
    ws.add_chart(ch, anc)
    return ch

barras(pl, 'A73', 'Número médio de pontos da trajetória por planeador',
       'Planeador', 'Pontos da trajetória',
       Reference(pl, min_col=2, max_col=3, min_row=67, max_row=70),
       Reference(pl, min_col=1, min_row=68, max_row=70), [C1, C2], w=15)

barras(pl, 'L63', 'Tempo médio de planeamento por planeador — ciclo completo, sem obstáculos',
       'Planeador', 'Tempo de planeamento (s)',
       Reference(pl, min_col=4, min_row=5, max_row=7),
       Reference(pl, min_col=1, min_row=5, max_row=7), [C1],
       err=['Planeadores!$E$5:$E$7'], w=15, leg=False)

barras(pl, 'L83', 'Tempo médio de planeamento por planeador — movimento com obstáculos',
       'Planeador', 'Tempo de planeamento (s)',
       Reference(pl, min_col=4, min_row=12, max_row=14),
       Reference(pl, min_col=1, min_row=12, max_row=14), [C2],
       err=['Planeadores!$E$12:$E$14'], w=15, leg=False)

# --- corrigir o gráfico das etapas de movimento -----------------------
# a série C61:C68 misturava milissegundos (respostas) com segundos (movimentos)
# no mesmo eixo. Passa a representar apenas as etapas medidas em segundos.
pp = wb['Pick-and-place']
ch = acha(pp, 'Tempo médio das etapas de movimento')
ch.title = 'Tempo médio das etapas de movimento (s)'
ref = Reference(pp, min_col=3, min_row=63, max_row=68)
ch.series[0].val.numRef.f = f"'Pick-and-place'!$C$63:$C$68"
ch.series[0].cat.strRef.f = f"'Pick-and-place'!$A$63:$A$68"
ch.y_axis.scaling.max = 16   # folga para as etiquetas de valor
pp['A70'] = ('Os tempos de resposta de PICK_POSE e de TARGET_POSE estão em milissegundos '
             'e por isso não figuram neste gráfico; constam da tabela acima.')
pp['A70'].font = Font(italic=True, size=9, color='595959')

# --- títulos que o LibreOffice não consegue compor sem cortar ----------
RENOMEAR = {
    'Tempo por etapa — cenário R01 (sem obstáculo)': 'Tempo por etapa — cenário R01',
    'Tempo por etapa — cenário O07 (com obstáculo)': 'Tempo por etapa — cenário O07',
    'Taxa de sucesso por comando e por ciclo': 'Sucesso por comando e por ciclo',
    'Ensaios de segurança e robustez por categoria': 'Ensaios de segurança e robustez',
}
for w in wb.worksheets:
    for c in w._charts:
        t = c.title.tx.rich.p[0].r[0].t if c.title else ''
        if t in RENOMEAR: c.title = RENOMEAR[t]

# --- sem legenda quando há uma só série -------------------------------
for w in wb.worksheets:
    for c in w._charts:
        if len(c.series) == 1:
            c.legend = None

# eixo vertical a começar em zero: num gráfico de barras, o comprimento da
# barra tem de ser proporcional ao valor. A exceção é o gráfico de diferenças,
# onde os valores são negativos.
for w in wb.worksheets:
    for ch in w._charts:
        t = ch.title.tx.rich.p[0].r[0].t if ch.title else ''
        if 'Diferença' not in t:
            ch.y_axis.scaling.min = 0

# dimensão uniforme: a que o LibreOffice compõe sem sobreposições
EMU = 360000
for w in wb.worksheets:
    for ch in w._charts:
        ch.width, ch.height = 17, 9.5
        if getattr(ch.anchor, 'ext', None) is not None:
            ch.anchor.ext.cx = int(17 * EMU); ch.anchor.ext.cy = int(9.5 * EMU)

wb.calculation.fullCalcOnLoad = True
wb.save('Graficos_Capitulo5_EDITAVEIS.xlsx')
print('gráficos totais:', sum(len(w._charts) for w in wb.worksheets))
