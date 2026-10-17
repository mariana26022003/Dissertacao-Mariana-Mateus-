# -*- coding: utf-8 -*-
"""Atualiza o livro de gráficos do Capítulo 5 para a base auditada.

Reescreve, a partir de Testes_de_tempo_e_estados_TM12_AUDITADO.xlsx:
  «Testes estatísticos»   pares concluídos com sucesso pelas duas estratégias, com as fórmulas realinhadas
  «Sincronização»         valores por cenário e agregados; títulos dos gráficos
  «Estatística descritiva» intervalos dos blocos por cenário e a tabela de Hodges-Lehmann
  «Testes emparelhados»   tabela de McNemar sobre os pares completos
  «Hodges-Lehmann»        grelhas de Walsh reconstruídas para os novos n

Revisão de 2026-10-08: com a exclusão das repetições 18 a 20 do O17 (sem registo
primário da estratégia temporal), a base passa a 147 condições e 138 pares. As
dimensões deixam de estar escritas no código: são lidas da base auditada, e a
última linha de dados do livro de origem é detetada em vez de assumida.

As folhas dos planeadores, do pick-and-place, do TCP, dos obstáculos e da
segurança não são tocadas: essas campanhas correram pela interface Qt3D e não
dependem do registo primário da supervisão. Os gráficos nativos são preservados.
"""
import copy, re, sys, numpy as np, openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from scipy import stats

AUD = 'Testes_de_tempo_e_estados_TM12_AUDITADO.xlsx'
SRC = sys.argv[1] if len(sys.argv) > 1 else 'Graficos_Capitulo5_EDITAVEIS_ORIGINAL.xlsx'
DST = sys.argv[2] if len(sys.argv) > 2 else 'Graficos_Capitulo5_EDITAVEIS.xlsx'
ORDEM = ['R01', 'R04', 'R08', 'R10', 'O07', 'O03', 'O13', 'O17']   # ordem da folha
TE = "'Testes estatísticos'"

# ───────────────────────────────────────────── 1. ler a base auditada
ws = openpyxl.load_workbook(AUD, data_only=True)['Dados Análise']
linhas = [[ws.cell(r, c).value for c in range(1, 61)] for r in range(2, ws.max_row + 1)]
linhas = [r for r in linhas if r[0] in ('Repetibilidade', 'Obstáculos')]
inc = [r for r in linhas if str(r[57]).strip().lower() == 'incluído']
par = [r for r in inc if r[7] == 1 and r[8] == 1]
assert par, 'nenhum par lido: o livro auditado tem os valores em cache? (abrir e gravar no Excel)'
print(f'condições incluídas: {len(inc)}   pares com sucesso nas duas estratégias: {len(par)}')

por_cen = {c: [] for c in ORDEM}
for r in par: por_cen[str(r[1]).strip()].append(r)
inc_cen = {c: [] for c in ORDEM}
for r in inc: inc_cen[str(r[1]).strip()].append(r)

# faixas de linhas na folha «Testes estatísticos», a partir da linha 6
faixas, lin = {}, 6
for c in ORDEM:
    n = len(por_cen[c]); faixas[c] = (lin, lin + n - 1); lin += n
FIM = lin - 1
assert FIM - 5 == len(par)
print('faixas:', {c: f'{a}-{b}' for c, (a, b) in faixas.items()}, '| última linha', FIM)

wb = openpyxl.load_workbook(SRC)

# ───────────────────────────────────────────── 2. «Testes estatísticos»
t = wb['Testes estatísticos']

FIM_ANT = max(r for r in range(6, t.max_row + 1) if t.cell(r, 1).value in ORDEM)
modelo = {c: copy.copy(t.cell(6, c)._style) for c in range(1, 8)}
for c in ORDEM:
    a, _ = faixas[c]
    for i, r in enumerate(por_cen[c]):
        rr = a + i
        t.cell(rr, 1, c)
        t.cell(rr, 2, float(r[3])).number_format = '0.000'
        t.cell(rr, 3, float(r[4])).number_format = '0.000'
        t.cell(rr, 4, f'=C{rr}-B{rr}').number_format = '0.000'
        t.cell(rr, 5, f'=ABS(D{rr})').number_format = '0.000'
        t.cell(rr, 6, f'=_xlfn.RANK.AVG(E{rr},$E$6:$E${FIM},1)').number_format = '0.0'
        t.cell(rr, 7, f'=IF(D{rr}<0,-1,1)').number_format = '0'
        for cc in range(1, 8): t.cell(rr, cc)._style = modelo[cc]
for rr in range(FIM + 1, FIM_ANT + 1):                # limpar o excedente
    for cc in range(1, 8): t.cell(rr, cc).value = None

# realinhar as referências que terminavam na antiga última linha (ex.: D6:D146 -> D6:D143)
def realinha(folha, ant, novo, so_te=False):
    k = 0
    for row in folha.iter_rows():
        for cel in row:
            v = cel.value
            if isinstance(v, str) and v.startswith('='):
                if so_te and 'Testes estatísticos' not in v: continue
                nv = re.sub(r'(\$?[A-Z]{1,2}\$?)' + str(ant) + r'\b', lambda m: m.group(1) + str(novo), v)
                if nv != v: cel.value = nv; k += 1
    return k
print('referências realinhadas:', realinha(t, FIM_ANT, FIM))

t['A4'] = f'1. DADOS EMPARELHADOS — os {len(par)} pares concluídos com sucesso pelas duas estratégias (Secção 5.4)'
t['J26'] = round(float(stats.shapiro(np.array([r[4] - r[3] for r in par], float)).pvalue), 5)
t['I27'] = ('Conclusão: normalidade rejeitada. O teste principal é o de Wilcoxon, definido por robustez e '
            'consistência entre métricas; o teste t fica como verificação.')
t['I21'] = 'Correlação rank-biserial (dimensão do efeito)'
n_ = len(par)
t['J21'] = f'=2*J10/({n_}*({n_}+1)/2)-1'
t['J21'].number_format = '0.000'

# secção 5 — McNemar substitui Fisher; atualizar as contagens de referência
suc_t = sum(1 for r in inc if r[7] == 1); suc_e = sum(1 for r in inc if r[8] == 1)
t['I30'] = (f'As duas estratégias executaram as mesmas {len(inc)} condições auditadas, pelo que os resultados '
            f'são emparelhados e o teste de Fisher, que pressupõe amostras independentes, não é aplicável. '
            f'O capítulo usa o teste exato de McNemar, na folha «Testes emparelhados» (p = 1,000). '
            f'O cálculo abaixo fica como registo do que foi feito antes da auditoria.')
t['J31'] = suc_t; t['K31'] = len(inc) - suc_t
t['J32'] = suc_e; t['K32'] = len(inc) - suc_e
# secção 6 — dimensões das amostras
t['I54'] = 'Pares emparelhados concluídos com sucesso (tempo de ciclo)'
t['J54'] = f'=COUNT(D6:D{FIM})'
t['I55'] = 'Condições auditadas por estratégia (taxa de sucesso)'; t['J55'] = len(inc)
t['K55'] = f'{len(linhas)} condições nominais − {len(linhas) - len(inc)} sem registo primário'
t['I56'] = 'Ciclos no total'; t['J56'] = 2 * len(inc); t['K56'] = f'{len(inc)} × 2 arquiteturas'

# ───────────────────────────────────────────── 3. «Sincronização»
s = wb['Sincronização']
s['A3'] = f'Dados por cenário (apenas os {len(par)} pares concluídos com sucesso pelas duas estratégias)'
for i, c in enumerate(ORDEM):
    r = 5 + i
    T = np.array([x[3] for x in por_cen[c]], float)
    E = np.array([x[4] for x in por_cen[c]], float)
    ET = np.array([x[5] for x in por_cen[c]], float)
    EE = np.array([x[6] for x in por_cen[c]], float)
    n = len(T); tc = stats.t.ppf(0.975, n - 1)
    fal_t = sum(1 for x in inc_cen[c] if x[7] != 1)
    fal_e = sum(1 for x in inc_cen[c] if x[8] != 1)
    vals = [c, n, T.mean(), E.mean(), tc * T.std(ddof=1) / np.sqrt(n), tc * E.std(ddof=1) / np.sqrt(n),
            ET.mean(), EE.mean(), fal_t, fal_e, 100 * (T.mean() - E.mean()) / T.mean()]
    for j, v in enumerate(vals):
        cel = s.cell(r, 1 + j, float(v) if isinstance(v, np.floating) else v)
        if j in (1, 8, 9): cel.number_format = '0'
        elif j > 1: cel.number_format = '0.000'

T = np.array([x[3] for x in par], float); E = np.array([x[4] for x in par], float)
ET = np.array([x[5] for x in par], float); EE = np.array([x[6] for x in par], float)
s['A57'] = f'Resultados agregados ({len(par)} pares para os tempos; {len(inc)} pares para a taxa de sucesso)'
for r, v_t, v_e in [(59, T.mean(), E.mean()), (60, np.median(T), np.median(E)),
                    (61, T.std(ddof=1), E.std(ddof=1)), (62, ET.mean(), EE.mean()),
                    (63, 100 * suc_t / len(inc), 100 * suc_e / len(inc))]:
    s.cell(r, 2, float(v_t)).number_format = '0.000'
    s.cell(r, 3, float(v_e)).number_format = '0.000'

# títulos com a dimensão da amostra: altera só o texto, preservando a formatação do título
TIT = [(re.compile(r'^Tempos agregados nas duas arquiteturas \(\d+ pares\)$'),
        f'Tempos agregados nas duas arquiteturas ({len(par)} pares)'),
       (re.compile(r'^Taxa de sucesso do ciclo \(\d+ ciclos por arquitetura\)$'),
        f'Taxa de sucesso do ciclo ({len(inc)} ciclos por arquitetura)')]
for w in wb.worksheets:
    for ch in w._charts:
        try: run = ch.title.tx.rich.p[0].r[0]
        except Exception: continue
        for padrao, novo in TIT:
            if padrao.match(run.t) and run.t != novo:
                print('título:', run.t, '->', novo); run.t = novo

# ───────────────────────────────────────────── 4. «Estatística descritiva»
d = wb['Estatística descritiva']
d['A4'] = f'1. TEMPO DE CICLO — os {len(par)} pares concluídos com sucesso pelas duas estratégias (Secção 5.4)'
COLF = {2: 'COUNT', 3: 'AVERAGE', 4: 'MEDIAN', 5: '_xlfn.STDEV.S', 6: 'MIN', 7: 'MAX'}
def bloco_agregado(r, col):
    rng = f'{TE}!{col}6:{col}{FIM}'
    for c, f in COLF.items(): d.cell(r, c, f'={f}({rng})')
    d.cell(r, 8, f'=QUARTILE({rng},1)'); d.cell(r, 9, f'=QUARTILE({rng},3)')
for r, col in [(6, 'B'), (7, 'C'), (8, 'D')]: bloco_agregado(r, col)
for base, col in [(14, 'B'), (26, 'C'), (38, 'D')]:
    for i, c in enumerate(ORDEM):
        a, b = faixas[c]; r = base + i; rng = f'{TE}!{col}{a}:{col}{b}'
        d.cell(r, 1, c)
        for cc, f in COLF.items(): d.cell(r, cc, f'={f}({rng})')
        d.cell(r, 8, f'=QUARTILE({rng},1)'); d.cell(r, 9, f'=QUARTILE({rng},3)')

def hodges_lehmann(x):
    x = np.asarray(x, float); n = len(x)
    w = np.sort(np.array([(x[i] + x[j]) / 2 for i in range(n) for j in range(i, n)]))
    k = int(round(n * (n + 1) / 4 - 1.959964 * np.sqrt(n * (n + 1) * (2 * n + 1) / 24)))
    return float(np.median(w)), float(w[k]), float(w[-k - 1])

HL = {}
for c in ORDEM:
    dd = np.array([x[4] - x[3] for x in por_cen[c]], float)
    HL[c] = (len(dd),) + hodges_lehmann(dd) + (float(stats.wilcoxon(dd).pvalue),)
dd = np.array([x[4] - x[3] for x in par], float)
HL['Agregado'] = (len(dd),) + hodges_lehmann(dd) + (float(stats.wilcoxon(dd).pvalue),)
for i, c in enumerate(ORDEM + ['Agregado']):
    r = 76 + i; n, hl, lo, hi, p = HL[c]
    X = par if c == 'Agregado' else por_cen[c]
    dif = [x[4] - x[3] for x in X]
    ic = stats.t.interval(0.95, n - 1, np.mean(dif), stats.sem(dif))
    for j, v in enumerate([c, n, hl, lo, hi, float(ic[0]), float(ic[1]), p]):
        cel = d.cell(r, 1 + j, v)
        if j == 1: cel.number_format = '0'
        elif j == 7: cel.number_format = '0.0000000'
        elif j > 1: cel.number_format = '0.000'
d['A86'] = ('Calculado por analise_estatistica_cap5.py e testes_emparelhados.py (scipy) sobre o livro '
            'auditado. O Excel não dispõe de função para as médias de Walsh na forma agregada, pelo que '
            'estes valores são reproduzíveis correndo os scripts sobre os mesmos ficheiros de ensaios. '
            'A folha «Hodges-Lehmann» reproduz o mesmo cálculo em fórmula.')

# ───────────────────────────────────────────── 5. «Testes emparelhados» — McNemar
e = wb['Testes emparelhados']
n11 = sum(1 for r in inc if r[7] == 1 and r[8] == 1)
n10 = sum(1 for r in inc if r[7] == 1 and r[8] != 1)
n01 = sum(1 for r in inc if r[7] != 1 and r[8] == 1)
n00 = sum(1 for r in inc if r[7] != 1 and r[8] != 1)
e['A176'] = (f'Cada repetição foi executada pelas duas estratégias nas mesmas condições, pelo que os '
             f'resultados são emparelhados e o teste de Fisher, que pressupõe amostras independentes, não é '
             f'aplicável. A tabela assenta nas {len(inc)} condições com registo primário das duas estratégias.')
e['B179'], e['C179'] = n11, n10
e['B180'], e['C180'] = n01, n00
print(f'McNemar: n11={n11} n10={n10} n01={n01} n00={n00} '
      f'p={stats.binomtest(n10, n10 + n01, 0.5).pvalue:.4f}')

# ───────────────────────────────────────────── 6. «Hodges-Lehmann» — grelhas de Walsh
h = wb['Hodges-Lehmann']
est_tit = copy.copy(h.cell(16, 1)._style)
for row in h.iter_rows(min_row=6):                     # limpar tudo abaixo do cabeçalho
    for cel in row: cel.value = None
h['A2'] = ('A mediana e os quartis descrevem a amostra. Para a diferença emparelhada é preciso um intervalo '
           'que não pressuponha normalidade, porque o teste de Shapiro-Wilk a rejeita. O método é o de '
           'Hodges-Lehmann: calculam-se as médias de todos os pares de diferenças — as médias de Walsh —, a '
           'estimativa é a mediana desse conjunto e os limites do intervalo são a k-ésima menor e a k-ésima '
           f'maior, com k dado pela distribuição de Wilcoxon. Base: os {len(par)} pares concluídos com '
           'sucesso pelas duas estratégias.')
lg, ancoras = 16, {}
for c in ORDEM + ['Agregado']:
    a, b = (6, FIM) if c == 'Agregado' else faixas[c]
    n = b - a + 1
    cel = h.cell(lg, 1, f'Médias de Walsh — {c} (n = {n})'); cel._style = est_tit
    r0 = lg + 1
    for i in range(n):
        for j in range(n):
            cl = h.cell(r0 + i, 2 + j, f'=IF({j+1}>={i+1},({TE}!$D${a+i}+{TE}!$D${a+j})/2,"")')
            cl.number_format = '0.000'
    ancoras[c] = (r0, 2, r0 + n - 1, 1 + n, n)
    lg = r0 + n + 2
for i, c in enumerate(ORDEM + ['Agregado']):
    r = 6 + i; r0, c0, r1, c1, n = ancoras[c]
    G = f'${get_column_letter(c0)}${r0}:${get_column_letter(c1)}${r1}'
    h.cell(r, 1, c).font = Font(size=10, bold=(c == 'Agregado'))
    h.cell(r, 2, n).number_format = '0'
    h.cell(r, 3, f'=ROUND(B{r}*(B{r}+1)/4-1.959964*SQRT(B{r}*(B{r}+1)*(2*B{r}+1)/24),0)+1').number_format = '0'
    h.cell(r, 4, f'=MEDIAN({G})').number_format = '0.000'
    h.cell(r, 5, f'=SMALL({G},C{r})').number_format = '0.000'
    h.cell(r, 6, f'=LARGE({G},C{r})').number_format = '0.000'
    h.cell(r, 7, f'=IF(OR(AND(E{r}<0,F{r}<0),AND(E{r}>0,F{r}>0)),"sim","não")')
h['A15'] = ('As grelhas que se seguem são o cálculo intermédio. Podem ser ocultadas sem afetar os '
            'resultados acima.')
h['A15'].font = Font(italic=True, size=9, color='595959')

wb.calculation.fullCalcOnLoad = True
wb.save(DST)
print('gravado:', DST)
print('gráficos preservados:', sum(len(w._charts) for w in wb.worksheets))
