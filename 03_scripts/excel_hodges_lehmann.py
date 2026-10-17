# -*- coding: utf-8 -*-
"""Acrescenta ao livro a folha «Hodges-Lehmann», com o cálculo completo em
fórmulas de Excel: médias de Walsh, estimativa da redução e intervalo de
confiança a 95 % distribution-free, sobre os pares concluídos com sucesso
pelas duas estratégias (138 na revisão de 2026-10-08). Reproduz exatamente a Tabela 29."""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from graficos_nativos import HDR2, BR, titulo

TE = "'Testes estatísticos'"

def faixas_da_folha(wb_):
    """Lê da folha «Testes estatísticos» os blocos contíguos de cada cenário (coluna A, a partir da
    linha 6). Evita escrever no código as linhas de cada cenário, que mudam com a base auditada."""
    t_ = wb_['Testes estatísticos']; blocos = []; r = 6
    while t_.cell(r, 1).value not in (None, ''):
        c = t_.cell(r, 1).value; a = r
        while t_.cell(r + 1, 1).value == c: r += 1
        blocos.append((c, a, r)); r += 1
    return blocos

wb = openpyxl.load_workbook('Graficos_Capitulo5_EDITAVEIS.xlsx')
# blocos contíguos das diferenças na folha «Testes estatísticos», coluna D
CEN = faixas_da_folha(wb)
AGREG = ('Agregado', 6, CEN[-1][2])
if 'Hodges-Lehmann' in wb.sheetnames: del wb['Hodges-Lehmann']
ws = wb.create_sheet('Hodges-Lehmann', 3)
ws.sheet_view.showGridLines = False
titulo(ws, 'A1', 'ESTIMATIVA DE HODGES-LEHMANN E INTERVALO DE CONFIANÇA A 95 % — TUDO EM FÓRMULA')

def nota(cel, t, alt=None):
    ws[cel] = t; ws[cel].font = Font(italic=True, size=9, color='595959')
    ws[cel].alignment = Alignment(wrap_text=True, vertical='top')
    if alt: ws.row_dimensions[int(cel[1:])].height = alt

nota('A2',
     'A mediana e os quartis descrevem a amostra. Para a diferença emparelhada é preciso um intervalo '
     'que não pressuponha normalidade, porque o teste de Shapiro-Wilk a rejeita. O método é o de '
     'Hodges-Lehmann: calculam-se as médias de todos os pares de diferenças — as médias de Walsh —, a '
     'estimativa é a mediana desse conjunto e os limites do intervalo são a k-ésima menor e a k-ésima '
     'maior, com k dado pela distribuição de Wilcoxon. É este o intervalo que corresponde ao teste de '
     'Wilcoxon, e por isso os dois nunca se contradizem.', 88)
nota('A3',
     'k = ARRED(n(n+1)/4 − 1,959964 × RAIZ(n(n+1)(2n+1)/24); 0) + 1. As grelhas triangulares abaixo '
     'contêm as médias de Walsh; as células vazias são texto e não entram em MED, MENOR nem MAIOR.', 34)

ws.column_dimensions['A'].width = 14
for c in 'BCDEFGH': ws.column_dimensions[c].width = 15

# ------------------------------------------------------------ tabela de resultados
CAB = ['Cenário', 'n', 'k', 'Redução HL (s)', 'IC 95 % inferior', 'IC 95 % superior', 'Exclui o zero?']
for j, h in enumerate(CAB):
    c = ws.cell(5, 1 + j, h); c.font = Font(bold=True, size=9.5); c.fill = HDR2
    c.alignment = Alignment(horizontal='center', wrap_text=True); c.border = BR

# ------------------------------------------------------------ grelhas de Walsh
linha_grelha = 16
ancoras = {}
for nome, a, b in CEN + [AGREG]:
    n = b - a + 1
    ws.cell(linha_grelha, 1, f'Médias de Walsh — {nome} (n = {n})').font = Font(bold=True, size=10, color='1F3864')
    r0 = linha_grelha + 1
    for i in range(n):
        for j in range(n):
            col = get_column_letter(2 + j)
            f = (f'=IF({j+1}>={i+1},({TE}!$D${a+i}+{TE}!$D${a+j})/2,"")')
            cel = ws.cell(r0 + i, 2 + j, f); cel.number_format = '0.000'
    ancoras[nome] = (r0, 2, r0 + n - 1, 1 + n, n)
    linha_grelha = r0 + n + 2

# ------------------------------------------------------------ fórmulas de resumo
for i, (nome, a, b) in enumerate(CEN + [AGREG]):
    r = 6 + i
    r0, c0, r1, c1, n = ancoras[nome]
    G = f'${get_column_letter(c0)}${r0}:${get_column_letter(c1)}${r1}'
    ws.cell(r, 1, nome).font = Font(size=10, bold=(nome == 'Agregado'))
    ws.cell(r, 2, n).number_format = '0'
    ws.cell(r, 3, f'=ROUND(B{r}*(B{r}+1)/4-1.959964*SQRT(B{r}*(B{r}+1)*(2*B{r}+1)/24),0)+1').number_format = '0'
    ws.cell(r, 4, f'=MEDIAN({G})').number_format = '0.000'
    ws.cell(r, 5, f'=SMALL({G},C{r})').number_format = '0.000'
    ws.cell(r, 6, f'=LARGE({G},C{r})').number_format = '0.000'
    ws.cell(r, 7, f'=IF(OR(AND(E{r}<0,F{r}<0),AND(E{r}>0,F{r}>0)),"sim","não")')
    for c in range(1, 8): ws.cell(r, c).border = BR

nota('A15', 'As grelhas que se seguem são o cálculo intermédio. Podem ser ocultadas sem afetar os '
            'resultados acima.')
wb.calculation.fullCalcOnLoad = True
wb.save('Graficos_Capitulo5_EDITAVEIS.xlsx')
print('folha «Hodges-Lehmann» criada; grelhas até à linha', linha_grelha)
