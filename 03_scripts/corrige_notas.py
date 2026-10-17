# -*- coding: utf-8 -*-
"""Alinha as notas metodológicas dos livros com o que os registos mostram.

O fluxo temporal não contém blocos de cancelamento. Nas condições incluídas na
análise comparativa não existe nenhum CANCEL nessa arquitetura. O único CANCEL
presente em temporal_ciclos.csv (TEMP_156) pertence a uma sessão posterior,
executada fora da sequência de ensaios, e não à repetição 20 do O17, como se
indicava na versão anterior desta nota (revisão de 2026-10-08).
"""
import openpyxl
from openpyxl.styles import Font

NOVA = ('Tempo de recuperação = intervalo entre a receção da primeira resposta de erro e a receção da '
        'resposta terminal ao cancelamento, medido no registo SentinelProcessEngine20260724 do VisionCore. '
        'A arquitetura temporal não possui ramo de erro: não lê a resposta, não deteta a falha e o fluxo '
        'não contém blocos de cancelamento. Nas 147 condições incluídas na análise comparativa não existe '
        'nenhum CANCEL nesta arquitetura. O único CANCEL presente no ficheiro temporal_ciclos.csv (TEMP_156) '
        'pertence a uma sessão posterior, executada fora da sequência de ensaios, e não integra a comparação. '
        'Os valores antes inscritos na coluna da arquitetura temporal são '
        'anotações manuais do tempo até à retoma pelo operador, não são comparáveis com os da arquitetura '
        'por respostas e não são reportados no Capítulo 5.')

wb = openpyxl.load_workbook('Testes_de_tempo_e_estados_TM12_AUDITADO.xlsx')
ws = wb['Pressupostos']
n = 0
for r in range(1, ws.max_row + 1):
    v = ws.cell(r, 1).value
    if isinstance(v, str) and v.startswith('Tempo de recuperação = intervalo'):
        ws.cell(r, 1, NOVA); n += 1
print('nota do livro auditado corrigida:', n)
wb.calculation.fullCalcOnLoad = True
wb.save('Testes_de_tempo_e_estados_TM12_AUDITADO.xlsx')

wb = openpyxl.load_workbook('Graficos_Capitulo5_EDITAVEIS.xlsx')
s = wb['Sincronização']
s['A64'] = 'Tempo médio de recuperação — apenas supervisão por respostas (s)'
s['A65'] = ('O tempo de recuperação abrange as seis ocorrências de resposta de erro registadas na supervisão '
            'por respostas; uma delas, a repetição 16 do cenário O03, pertence a uma condição excluída da '
            'comparação emparelhada por falta do registo primário da estratégia temporal. A supervisão '
            'temporal não possui grandeza equivalente: o fluxo não lê as respostas, não deteta a falha e não '
            'contém blocos de cancelamento, pelo que o gráfico correspondente foi removido.')
s['A65'].font = Font(italic=True, size=9, color='595959')
wb.calculation.fullCalcOnLoad = True
wb.save('Graficos_Capitulo5_EDITAVEIS.xlsx')
print('nota do livro de gráficos corrigida')
