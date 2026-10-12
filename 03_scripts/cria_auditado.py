# -*- coding: utf-8 -*-
"""Cria a cópia auditada do livro de tempos e estados.

Acrescenta à folha «Dados Análise» as colunas estado_inclusao, motivo_exclusao e
estrategia_sem_registo, e uma folha «Auditoria de inclusão» que documenta o
critério e as contagens. Nenhum valor original é alterado ou apagado.
"""
import shutil, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ORIG = 'Testes_de_tempo_e_estados_TM12_CORRIGIDO.xlsx'
DEST = 'Testes_de_tempo_e_estados_TM12_AUDITADO.xlsx'

# Condições sem registo primário, por estratégia (auditoria de proveniência)
SEM_REGISTO = {
    ('R10', 20): 'TEMPORAL',
    ('O03', 16): 'TEMPORAL', ('O03', 17): 'TEMPORAL', ('O03', 18): 'TEMPORAL',
    ('O03', 19): 'TEMPORAL', ('O03', 20): 'TEMPORAL',
    ('O13', 17): 'ESTADOS', ('O13', 18): 'ESTADOS',
    ('O13', 19): 'ESTADOS', ('O13', 20): 'ESTADOS',
    # revisão de 2026-10-08: no registo SentinelProcessEngine20260723, a repetição 18
    # do O17 (temporal) foi interrompida às 15:19:36 e as repetições 19 e 20 não foram
    # executadas; os ciclos TEMP_154 a TEMP_156 pertencem a uma sessão posterior.
    ('O17', 18): 'TEMPORAL', ('O17', 19): 'TEMPORAL', ('O17', 20): 'TEMPORAL',
}
MOTIVO = ('Ausência de registo primário da estratégia {e}; par excluído para '
          'preservar o desenho emparelhado.')
NOTA_O17 = ('Na estratégia temporal, o registo SentinelProcessEngine20260723 do VisionCore mostra dezassete '
            'ciclos completos do cenário O17, entre as 15:04 e as 15:18. A repetição 18 foi iniciada às '
            '15:18:55 e o fluxo foi interrompido às 15:19:36, durante a aproximação à peça, e as repetições '
            '19 e 20 não foram executadas nessa sessão. Os ciclos TEMP_154 a TEMP_156 do ficheiro '
            'temporal_ciclos.csv pertencem a uma sessão posterior, com outro processo e outras coordenadas, '
            'e não correspondem a estas condições. As três condições ficam por isso sem registo primário da '
            'estratégia temporal e são excluídas pelo mesmo critério das restantes. Os valores que constavam '
            'da base para estas condições mantêm-se na folha «Obs_Ciclos», sem alteração, mas não integram '
            'a análise.')

shutil.copyfile(ORIG, DEST)
wb = openpyxl.load_workbook(DEST)
ws = wb['Dados Análise']
wv = openpyxl.load_workbook(ORIG, data_only=True)['Dados Análise']   # valores em cache

HDR = PatternFill('solid', fgColor='DDEBF7')
BR = Border(*[Side('thin', color='BFBFBF')] * 4)
for j, t in enumerate(['estado_inclusao', 'motivo_exclusao', 'estrategia_sem_registo']):
    c = ws.cell(1, 58 + j, t)
    c.font = Font(bold=True, size=9.5); c.fill = HDR
    c.alignment = Alignment(horizontal='center', wrap_text=True); c.border = BR
ws.column_dimensions['BF'].width = 15
ws.column_dimensions['BG'].width = 52
ws.column_dimensions['BH'].width = 22

n_inc = n_exc = 0
for r in range(2, ws.max_row + 1):
    conj = ws.cell(r, 1).value
    if conj not in ('Repetibilidade', 'Obstáculos'):
        continue
    k = (str(ws.cell(r, 2).value).strip(), ws.cell(r, 3).value)
    if k in SEM_REGISTO:
        e = SEM_REGISTO[k]
        ws.cell(r, 58, 'excluído'); ws.cell(r, 59, MOTIVO.format(e=e)); ws.cell(r, 60, e)
        n_exc += 1
    else:
        ws.cell(r, 58, 'incluído'); n_inc += 1
    for c in range(58, 61): ws.cell(r, c).border = BR

# ---------------------------------------------------- folha de auditoria
if 'Auditoria de inclusão' in wb.sheetnames: del wb['Auditoria de inclusão']
a = wb.create_sheet('Auditoria de inclusão')
a.sheet_view.showGridLines = False
a.column_dimensions['A'].width = 46
for c in 'BCD': a.column_dimensions[c].width = 22
a['A1'] = 'AUDITORIA DE PROVENIÊNCIA E CRITÉRIO DE INCLUSÃO'
a['A1'].font = Font(bold=True, size=13, color='1F3864')
a['A2'] = ('Uma condição cenário-repetição só entra na análise comparativa quando existe registo primário '
           'das duas estratégias. Onde falta o registo de uma delas, a condição e o respetivo par são '
           'excluídos, para preservar o desenho emparelhado. A exclusão é determinada exclusivamente pela '
           'disponibilidade de evidência primária e nunca pelo resultado do ensaio. Nenhum valor original '
           'foi alterado ou apagado, e nenhuma observação foi imputada, estimada ou reconstruída.')
a['A2'].font = Font(italic=True, size=9.5, color='595959')
a['A2'].alignment = Alignment(wrap_text=True, vertical='top')
a.row_dimensions[2].height = 62

lin = 4
a.cell(lin, 1, 'Contagens').font = Font(bold=True, size=11, color='1F3864'); lin += 1
suc_t = suc_e = amb = 0
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 58).value != 'incluído': continue
    st, se = wv.cell(r, 8).value, wv.cell(r, 9).value
    suc_t += (st == 1); suc_e += (se == 1); amb += (st == 1 and se == 1)
for rot, val in [('Condições no desenho nominal, por estratégia', n_inc + n_exc),
                 ('Condições excluídas por falta de registo primário', n_exc),
                 ('Condições incluídas, por estratégia', n_inc),
                 ('Ciclos incluídos no total', 2 * n_inc),
                 ('Pares concluídos com sucesso pelas duas estratégias', amb),
                 ('Ciclos bem-sucedidos, supervisão temporal', suc_t),
                 ('Ciclos bem-sucedidos, supervisão por respostas', suc_e)]:
    a.cell(lin, 1, rot).font = Font(size=10)
    c = a.cell(lin, 2, val); c.font = Font(size=10, bold=True); c.number_format = '0'
    a.cell(lin, 1).border = BR; c.border = BR
    lin += 1

lin += 2
a.cell(lin, 1, 'Condições excluídas').font = Font(bold=True, size=11, color='1F3864'); lin += 1
for j, t in enumerate(['Cenário', 'Repetição', 'Estratégia sem registo', 'Motivo']):
    c = a.cell(lin, 1 + j, t); c.font = Font(bold=True, size=9.5); c.fill = HDR
    c.alignment = Alignment(horizontal='center'); c.border = BR
a.column_dimensions['D'].width = 60
lin += 1
for (cen, rep), e in sorted(SEM_REGISTO.items()):
    for j, v in enumerate([cen, rep, e, MOTIVO.format(e=e)]):
        c = a.cell(lin, 1 + j, v); c.font = Font(size=9.5); c.border = BR
        c.alignment = Alignment(wrap_text=(j == 3), vertical='top')
    lin += 1

lin += 1
a.cell(lin, 1, 'Nota sobre as repetições 18 a 20 do cenário O17 (revisão de 2026-10-08)').font = \
    Font(bold=True, size=11, color='1F3864')
a.cell(lin + 1, 1, NOTA_O17).font = Font(italic=True, size=9.5, color='595959')
a.cell(lin + 1, 1).alignment = Alignment(wrap_text=True, vertical='top')
a.merge_cells(start_row=lin + 1, start_column=1, end_row=lin + 1, end_column=4)
a.row_dimensions[lin + 1].height = 120

wb.calculation.fullCalcOnLoad = True
wb.save(DEST)
print(f'incluídas: {n_inc}   excluídas: {n_exc}   pares com sucesso em ambas: {amb}')
print(f'sucesso temporal: {suc_t}/{n_inc}   sucesso por respostas: {suc_e}/{n_inc}')
