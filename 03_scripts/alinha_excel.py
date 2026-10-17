# -*- coding: utf-8 -*-
"""Alinha o livro de gráficos com a versão final do Capítulo 5:
assinala os testes substituídos e acrescenta os intervalos de Hodges-Lehmann."""
import numpy as np, openpyxl
from scipy import stats
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from graficos_nativos import HDR2, BR

AVISO = PatternFill('solid', fgColor='FFF3CD')

# base auditada: condições incluídas e pares concluídos com sucesso pelas duas estratégias
_a = openpyxl.load_workbook('Testes_de_tempo_e_estados_TM12_AUDITADO.xlsx', data_only=True)['Dados Análise']
_L = [[_a.cell(r, c).value for c in range(1, 61)] for r in range(2, _a.max_row + 1)]
INC = [r for r in _L if r[0] in ('Repetibilidade', 'Obstáculos') and str(r[57]).strip().lower() == 'incluído']
PAR = [r for r in INC if r[7] == 1 and r[8] == 1]
wb = openpyxl.load_workbook('Graficos_Capitulo5_EDITAVEIS.xlsx')

def marca(ws, cel, txt):
    ws[cel] = txt
    ws[cel].font = Font(bold=True, size=9.5, color='8A6100')
    ws[cel].fill = AVISO
    ws[cel].alignment = Alignment(wrap_text=True, vertical='center')

# --------------------------------------------- 1. Fisher deixou de ser aplicável
ws = wb['Testes estatísticos']
alvo = None
for r in range(1, ws.max_row + 1):
    v = ws.cell(r, 9).value
    if isinstance(v, str) and v.startswith('5. TESTE EXATO DE FISHER'):
        alvo = r; break
if alvo:
    ws.cell(alvo, 9, '5. TESTE EXATO DE FISHER — SUBSTITUÍDO, ver folha «Testes emparelhados»')
    ws.cell(alvo, 9).font = Font(bold=True, size=10.5, color='8A6100')
    marca(ws, f'I{alvo+1}',
          f'As duas estratégias executaram as mesmas {len(INC)} condições auditadas, pelo que os resultados '
          'são emparelhados e o teste de Fisher, que pressupõe amostras independentes, não é aplicável. '
          'O capítulo usa o teste exato de McNemar, calculado na folha «Testes emparelhados» '
          '(p = 1,000). O cálculo abaixo fica como registo do que foi feito antes da auditoria.')

# ------------------------------- 2. Kruskal-Wallis e Mann-Whitney substituídos
ws = wb['Dados planeadores']
for r in range(1, ws.max_row + 1):
    v = ws.cell(r, 8).value
    if isinstance(v, str) and v.startswith(('KRUSKAL-WALLIS', 'MANN-WHITNEY')):
        ws.cell(r, 8, v + ' — SUBSTITUÍDO')
        ws.cell(r, 8).font = Font(bold=True, size=10.5, color='8A6100')
marca(ws, 'H1',
      'Os três planeadores executaram exatamente os mesmos movimentos: o desenho é em blocos '
      'emparelhados e os testes para grupos independentes não são aplicáveis. O capítulo usa '
      'Friedman, Wilcoxon com correção de Holm e Q de Cochran, na folha «Testes emparelhados». '
      'Os cálculos desta folha ficam como registo da versão anterior.')
ws.column_dimensions['H'].width = 46
ws.row_dimensions[1].height = 60

# ------------------------------------- 3. Hodges-Lehmann na estatística descritiva
ws = wb['Estatística descritiva']
r = ws.max_row + 3
ws.cell(r, 1, '8. REDUÇÃO DO TEMPO DE CICLO — ESTIMATIVA DE HODGES-LEHMANN')
ws.cell(r, 1).font = Font(bold=True, size=10.5, color='1F3864')
marca(ws, f'A{r+1}',
      'Os intervalos dos blocos 1 a 6 são calculados pela distribuição t e pressupõem normalidade. '
      'Como o teste de Shapiro-Wilk a rejeita para as diferenças emparelhadas, o Capítulo 5 reporta '
      'para essas diferenças a estimativa de Hodges-Lehmann e o intervalo distribution-free '
      'correspondente ao teste de Wilcoxon, que é o que consta da Tabela 29. Os dois métodos '
      'coincidem no valor agregado; divergem no cenário R01, onde o intervalo pela t exclui o zero '
      'por dois milésimos e o de Hodges-Lehmann não o exclui, em coerência com p = 0,067. Os valores '
      f'assentam nos {len(PAR)} pares concluídos com sucesso pelas duas estratégias.')
ws.row_dimensions[r+1].height = 74

HDR = ['Cenário', 'n', 'Redução HL (s)', 'IC 95 % HL inferior', 'IC 95 % HL superior',
       'IC 95 % t inferior', 'IC 95 % t superior', 'p de Wilcoxon']
for j, h in enumerate(HDR):
    c = ws.cell(r+3, 1+j, h); c.font = Font(bold=True, size=9.5); c.fill = HDR2
    c.alignment = Alignment(horizontal='center', wrap_text=True); c.border = BR
# valores calculados sobre a base auditada; ver analise_estatistica_cap5.py
def _hl(x):
    x = np.asarray(x, float); n = len(x)
    w = np.sort([(x[i] + x[j]) / 2 for i in range(n) for j in range(i, n)])
    k = int(round(n * (n + 1) / 4 - 1.959964 * np.sqrt(n * (n + 1) * (2 * n + 1) / 24)))
    return float(np.median(w)), float(w[k]), float(w[-k - 1])
DADOS = []
for _c in ['R01', 'R04', 'R08', 'R10', 'O07', 'O03', 'O13', 'O17', 'Agregado']:
    _X = PAR if _c == 'Agregado' else [r for r in PAR if str(r[1]).strip() == _c]
    _d = np.array([r[4] - r[3] for r in _X], float)
    _h = _hl(_d); _ic = stats.t.interval(0.95, len(_d) - 1, _d.mean(), stats.sem(_d))
    DADOS.append((_c, len(_d), _h[0], _h[1], _h[2], float(_ic[0]), float(_ic[1]),
                  float(stats.wilcoxon(_d).pvalue)))
for i, linha in enumerate(DADOS):
    for j, v in enumerate(linha):
        c = ws.cell(r+4+i, 1+j, v); c.font = Font(size=10); c.border = BR
        if isinstance(v, float): c.number_format = '0.000' if j < 7 else '0.0000000'
        if isinstance(v, int) and j == 1: c.number_format = '0'
fim = r + 4 + len(DADOS)
ws.cell(fim+1, 1, 'Calculado por testes_emparelhados.py e analise_estatistica_cap5.py (scipy). O Excel não '
                  'dispõe de função para as médias de Walsh, pelo que estes valores são reproduzíveis '
                  'correndo os scripts sobre os mesmos ficheiros de ensaios.')
ws.cell(fim+1, 1).font = Font(italic=True, size=9, color='595959')

wb.calculation.fullCalcOnLoad = True
wb.save('Graficos_Capitulo5_EDITAVEIS.xlsx')
print('alinhado. folhas:', len(wb.sheetnames))
