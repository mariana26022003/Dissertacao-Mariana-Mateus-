# -*- coding: utf-8 -*-
"""
Análise de sensibilidade da comparação entre estratégias
========================================================
Repete a comparação do tempo de ciclo e da espera improdutiva apenas sobre os
pares cuja correspondência ao registo primário foi confirmada diretamente nas
duas estratégias (Tabelas S8 e S9 do Online Resource 1).

Critério: dos 138 pares concluídos com sucesso pelas duas estratégias, entram
os pares em que a coluna «Estado da identificação» da tabela de correspondência
é «confirmado» tanto para a estratégia TEMPORAL como para a estratégia ESTADOS.
Os restantes pares foram identificados pela posição na sequência cronológica
(«inferido») e ficam de fora desta análise.

Execução, a partir da pasta raiz do repositório:
    python 03_scripts/analise_sensibilidade.py

Ficheiros necessários:
    Testes_de_tempo_e_estados_TM12_AUDITADO.xlsx
    Tabela_Correspondencia_Ciclos_AUDITADA.xlsx
"""
import numpy as np
import openpyxl
from scipy import stats

AUD = 'Testes_de_tempo_e_estados_TM12_AUDITADO.xlsx'
CORR = 'Tabela_Correspondencia_Ciclos_AUDITADA.xlsx'


def hodges_lehmann(d):
    """Estimador de Hodges-Lehmann e intervalo de confiança a 95 % livre de distribuição."""
    d = np.asarray(d, float); n = len(d)
    w = np.sort([(d[i] + d[j]) / 2 for i in range(n) for j in range(i, n)])
    k = int(round(n * (n + 1) / 4 - 1.959964 * np.sqrt(n * (n + 1) * (2 * n + 1) / 24)))
    return float(np.median(w)), float(w[k]), float(w[-k - 1])


def descrever(rotulo, x):
    x = np.asarray(x, float)
    q1, q3 = np.percentile(x, [25, 75])
    ic = stats.t.interval(0.95, len(x) - 1, x.mean(), stats.sem(x))
    print(f'    {rotulo:<20} média {x.mean():8.3f}   mediana {np.median(x):8.3f}   DP {x.std(ddof=1):6.3f}   '
          f'mín. {x.min():8.3f}   máx. {x.max():8.3f}   IQR {q3 - q1:6.3f}   IC 95 % [{ic[0]:.3f}; {ic[1]:.3f}]')


def comparar(nome, temporal, respostas):
    temporal = np.asarray(temporal, float); respostas = np.asarray(respostas, float)
    d = respostas - temporal
    n = len(d)
    print(f'\n-- {nome} (n = {n} pares) --')
    descrever('temporal', temporal)
    descrever('por respostas', respostas)
    descrever('diferença', d)
    print(f'    variação relativa da média: {100 * (respostas.mean() - temporal.mean()) / temporal.mean():+.2f} %')
    sw = stats.shapiro(d)
    print(f'    Shapiro-Wilk sobre as diferenças: W = {sw.statistic:.3f}   p = {sw.pvalue:.5f}'
          f'   -> normalidade {"não rejeitada" if sw.pvalue > 0.05 else "rejeitada"}')
    w = stats.wilcoxon(d)
    print(f'    Wilcoxon emparelhado (teste principal): V = {w.statistic:.0f}   p = {w.pvalue:.3e}')
    tt = stats.ttest_rel(respostas, temporal)
    print(f'    teste t emparelhado (verificação): t({n - 1}) = {tt.statistic:.3f}   p = {tt.pvalue:.3e}')
    hl, lo, hi = hodges_lehmann(d)
    print(f'    Hodges-Lehmann = {hl:.3f}   IC 95 % = [{lo:.3f}; {hi:.3f}]')
    postos = stats.rankdata(np.abs(d))
    rb = (postos[d > 0].sum() - postos[d < 0].sum()) / (n * (n + 1) / 2)
    print(f'    correlação rank-biserial = {rb:+.3f}')
    print(f'    diferenças nulas = {int((d == 0).sum())}   valores absolutos repetidos = '
          f'{n - len(np.unique(np.abs(d)))}')


# ---------------------------------------------------------------- dados
ws = openpyxl.load_workbook(AUD, data_only=True)['Dados Análise']
linhas = [[ws.cell(r, c).value for c in range(1, 61)] for r in range(2, ws.max_row + 1)]
linhas = [r for r in linhas if r[0] in ('Repetibilidade', 'Obstáculos')]
incluidas = [r for r in linhas if str(r[57]).strip().lower() == 'incluído']
pares = [r for r in incluidas if r[7] == 1 and r[8] == 1]

cw = openpyxl.load_workbook(CORR, data_only=True)['Correspondência']
estado = {}
for r in cw.iter_rows(min_row=2, values_only=True):
    if r[4]:   # identificador do ciclo
        estado[(str(r[1]).strip(), r[2], r[3])] = r[10]   # (cenário, repetição, estratégia) -> estado

def confirmado(r):
    k = (str(r[1]).strip(), r[2])
    return estado.get(k + ('TEMPORAL',)) == 'confirmado' and estado.get(k + ('ESTADOS',)) == 'confirmado'

conf = [r for r in pares if confirmado(r)]
recon = [r for r in pares if not confirmado(r)]

print('ANÁLISE DE SENSIBILIDADE — pares com correspondência confirmada nas duas estratégias')
print('=' * 78)
print(f'Condições incluídas: {len(incluidas)}   pares concluídos pelas duas estratégias: {len(pares)}')
print(f'Pares com correspondência confirmada nas duas estratégias: {len(conf)}')
print(f'Pares com correspondência reconstruída em pelo menos uma estratégia: {len(recon)}')
for r in recon:
    k = (str(r[1]).strip(), r[2])
    print(f'    {k[0]} repetição {k[1]:>2}   temporal: {estado.get(k + ("TEMPORAL",))}   '
          f'por respostas: {estado.get(k + ("ESTADOS",))}')

comparar('Tempo total de ciclo (s)', [r[3] for r in conf], [r[4] for r in conf])
comparar('Espera improdutiva por ciclo (s)', [r[5] for r in conf], [r[6] for r in conf])
print('\n' + '=' * 78)
print('Fim. Os valores correspondem às Tabelas S8 e S9 e ao parágrafo que as precede no Online Resource 1.')
