# -*- coding: utf-8 -*-
"""Acrescenta à tabela de correspondência o estado de inclusão de cada ciclo,
lido da base auditada. Não altera nenhuma coluna existente."""
import shutil, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

AUD = 'Testes_de_tempo_e_estados_TM12_AUDITADO.xlsx'
ORIG = 'Tabela_Correspondencia_Ciclos.xlsx'
DEST = 'Tabela_Correspondencia_Ciclos_AUDITADA.xlsx'

a = openpyxl.load_workbook(AUD, data_only=True)['Dados Análise']
estado = {}
for r in range(2, a.max_row + 1):
    if a.cell(r, 1).value not in ('Repetibilidade', 'Obstáculos'): continue
    estado[(str(a.cell(r, 2).value).strip(), a.cell(r, 3).value)] = (
        a.cell(r, 58).value, a.cell(r, 59).value, a.cell(r, 60).value)

# Detalhe da exclusão do O17 18-20 (revisão de 2026-10-08), escrito na tabela de correspondência
# porque é aqui que estão os identificadores TEMP_154 a TEMP_156 antes atribuídos a estas condições.
DETALHE_O17 = ('Ausência de registo primário da estratégia temporal; par excluído para preservar o desenho '
               'emparelhado. No registo SentinelProcessEngine20260723, a repetição 18 do O17 foi interrompida '
               'às 15:19:36 e as repetições 19 e 20 não foram executadas; os ciclos TEMP_154 a TEMP_156 '
               'pertencem a uma sessão posterior e não correspondem a estas condições.')
for k in [('O17', 18), ('O17', 19), ('O17', 20)]:
    if k in estado and estado[k][0] == 'excluído':
        estado[k] = (estado[k][0], DETALHE_O17, estado[k][2])

shutil.copyfile(ORIG, DEST)
wb = openpyxl.load_workbook(DEST)
ws = wb['Correspondência']
HDR = PatternFill('solid', fgColor='DDEBF7'); BR = Border(*[Side('thin', color='BFBFBF')] * 4)
c0 = ws.max_column + 1
for j, t in enumerate(['estado_inclusao', 'motivo_exclusao', 'estrategia_sem_registo']):
    c = ws.cell(1, c0 + j, t); c.font = Font(bold=True, size=9.5); c.fill = HDR
    c.alignment = Alignment(horizontal='center', wrap_text=True); c.border = BR
ws.column_dimensions[ws.cell(1, c0 + 1).column_letter].width = 52

n = {'incluído': 0, 'excluído': 0}
for r in range(2, ws.max_row + 1):
    if not ws.cell(r, 5).value: continue
    k = (str(ws.cell(r, 2).value).strip(), ws.cell(r, 3).value)
    v = estado.get(k)
    if not v: continue
    for j in range(3):
        cel = ws.cell(r, c0 + j, v[j]); cel.border = BR
        cel.alignment = Alignment(wrap_text=(j == 1), vertical='top')
    n[v[0]] = n.get(v[0], 0) + 1

# resumo
rs = wb['Resumo']
lin = rs.max_row + 3
rs.cell(lin, 1, 'Estado de inclusão na análise comparativa').font = Font(bold=True, size=11, color='1F3864')
for i, (rot, val) in enumerate([('Ciclos incluídos', n['incluído']), ('Ciclos excluídos', n['excluído'])], 1):
    rs.cell(lin + i, 1, rot).font = Font(size=10)
    c = rs.cell(lin + i, 2, val); c.font = Font(size=10, bold=True); c.number_format = '0'
rs.cell(lin + 3, 1, 'Critério: uma condição cenário-repetição só é incluída quando existe registo primário '
                    'das duas estratégias. Nenhum valor original foi alterado ou apagado.').font = \
    Font(italic=True, size=9, color='595959')

wb.calculation.fullCalcOnLoad = True
wb.save(DEST)
print('ciclos anotados:', n)
